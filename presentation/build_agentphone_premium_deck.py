from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT_DIR = Path(__file__).resolve().parent
OUT_FILE = OUT_DIR / "AgentPhone_Premium_Sales_Deck.pptx"

W = Inches(13.333)
H = Inches(7.5)

INK = RGBColor(7, 10, 18)
INK_2 = RGBColor(11, 16, 28)
CARD = RGBColor(19, 28, 48)
CARD_2 = RGBColor(25, 38, 63)
WHITE = RGBColor(248, 250, 252)
MIST = RGBColor(171, 184, 205)
DIM = RGBColor(119, 135, 162)
MINT = RGBColor(80, 240, 178)
BLUE = RGBColor(92, 154, 255)
GOLD = RGBColor(255, 202, 103)
ROSE = RGBColor(255, 116, 148)
VIOLET = RGBColor(171, 130, 255)

FONT = "Aptos"


def rgb(hex_value):
    value = hex_value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def add_text(slide, x, y, w, h, text, size=18, color=WHITE, bold=False, align=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bg(slide, mood="mint"):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = INK

    colors = {
        "mint": (rgb("#123E3A"), rgb("#182448")),
        "blue": (rgb("#142B57"), rgb("#112F35")),
        "rose": (rgb("#402033"), rgb("#132342")),
        "gold": (rgb("#3E3217"), rgb("#142942")),
        "violet": (rgb("#2D245C"), rgb("#113740")),
    }
    c1, c2 = colors[mood]
    for x, y, size, color, transparency in [
        (Inches(9.2), Inches(-1.1), Inches(5.0), c1, 25),
        (Inches(-1.4), Inches(5.0), Inches(4.6), c2, 35),
        (Inches(5.4), Inches(1.8), Inches(2.4), rgb("#1A2740"), 58),
    ]:
        glow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x, y, size, size)
        glow.fill.solid()
        glow.fill.fore_color.rgb = color
        glow.fill.transparency = transparency
        glow.line.fill.background()


def add_card(slide, x, y, w, h, fill=CARD, line=rgb("#35445F"), transparency=0):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    shape.line.color.rgb = line
    shape.line.transparency = 30
    return shape


def add_pill(slide, x, y, text, color=MINT, width=1.8):
    pill = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, Inches(width), Inches(0.34))
    pill.fill.solid()
    pill.fill.fore_color.rgb = color
    pill.line.fill.background()
    tf = pill.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = INK
    return pill


def add_title(slide, kicker, title, subtitle=None, mood_color=MINT):
    add_pill(slide, Inches(0.72), Inches(0.46), kicker.upper(), mood_color, 2.05)
    add_text(slide, Inches(0.72), Inches(0.98), Inches(11.7), Inches(0.78), title, 34, WHITE, True)
    if subtitle:
        add_text(slide, Inches(0.75), Inches(1.76), Inches(10.9), Inches(0.5), subtitle, 15.5, MIST)


def add_footer(slide, index):
    add_text(slide, Inches(0.72), Inches(7.05), Inches(2.2), Inches(0.18), "AgentPhone / premium v2", 8.5, DIM, True)
    add_text(slide, Inches(12.0), Inches(7.05), Inches(0.6), Inches(0.18), f"{index:02}", 8.5, DIM, True, PP_ALIGN.RIGHT)


def add_bullets(slide, x, y, w, h, items, size=15, color=WHITE, gap=0.1):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(gap * 72)
        p._p.get_or_add_pPr().set("marL", "0")
    return box


def add_metric(slide, x, y, number, label, color=MINT):
    add_card(slide, x, y, Inches(2.35), Inches(1.08), CARD_2)
    add_text(slide, x + Inches(0.22), y + Inches(0.17), Inches(1.9), Inches(0.36), number, 26, color, True)
    add_text(slide, x + Inches(0.24), y + Inches(0.68), Inches(1.85), Inches(0.26), label, 10.5, MIST)


def add_phone(slide, x, y, w, h, accent=MINT):
    shell = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shell.fill.solid()
    shell.fill.fore_color.rgb = rgb("#050812")
    shell.line.color.rgb = rgb("#647899")
    shell.line.width = Pt(1.1)

    screen = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x + Inches(0.14), y + Inches(0.2), w - Inches(0.28), h - Inches(0.4))
    screen.fill.solid()
    screen.fill.fore_color.rgb = INK_2
    screen.line.fill.background()

    add_text(slide, x + Inches(0.34), y + Inches(0.45), w - Inches(0.68), Inches(0.22), "AgentPhone", 10.5, MIST, True, PP_ALIGN.CENTER)
    add_bubble(slide, x + Inches(0.35), y + Inches(1.08), w - Inches(0.7), "Согласуй вечер с Васей", False, accent)
    add_bubble(slide, x + Inches(0.35), y + Inches(1.84), w - Inches(0.7), "Нашёл 3 варианта. Написать и забронировать?", True, accent)
    add_bubble(slide, x + Inches(0.35), y + Inches(2.88), w - Inches(0.7), "Да", False, accent)
    add_bubble(slide, x + Inches(0.35), y + Inches(3.56), w - Inches(0.7), "Готово. Я напомню, когда пора выходить.", True, accent)
    mic = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x + w / 2 - Inches(0.3), y + h - Inches(0.82), Inches(0.6), Inches(0.6))
    mic.fill.solid()
    mic.fill.fore_color.rgb = accent
    mic.line.fill.background()
    add_text(slide, x + w / 2 - Inches(0.07), y + h - Inches(0.65), Inches(0.14), Inches(0.12), "•", 19, INK, True, PP_ALIGN.CENTER)


def add_bubble(slide, x, y, w, text, agent, accent):
    bubble_w = w * 0.82
    bx = x if agent else x + w - bubble_w
    fill = rgb("#21314F") if agent else accent
    color = WHITE if agent else INK
    bubble = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, bx, y, bubble_w, Inches(0.58))
    bubble.fill.solid()
    bubble.fill.fore_color.rgb = fill
    bubble.line.fill.background()
    tf = bubble.text_frame
    tf.clear()
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(9)
    r.font.color.rgb = color


def add_case_card(slide, x, y, number, title, body, color):
    add_card(slide, x, y, Inches(2.25), Inches(2.02), CARD_2)
    add_text(slide, x + Inches(0.2), y + Inches(0.18), Inches(0.46), Inches(0.28), number, 19, color, True)
    add_text(slide, x + Inches(0.68), y + Inches(0.2), Inches(1.26), Inches(0.35), title, 14.5, WHITE, True)
    add_text(slide, x + Inches(0.22), y + Inches(0.82), Inches(1.82), Inches(0.78), body, 10.7, MIST)


def add_case_slide(prs, idx, mood, accent, case_num, title, promise, user_says, actions, result):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, mood)
    add_title(slide, f"Pocket case {case_num}", title, promise, accent)

    add_card(slide, Inches(0.86), Inches(2.58), Inches(4.0), Inches(3.02), CARD_2)
    add_text(slide, Inches(1.2), Inches(2.96), Inches(3.25), Inches(0.35), "Пользователь говорит", 16, accent, True)
    add_text(slide, Inches(1.22), Inches(3.62), Inches(3.1), Inches(0.82), f"«{user_says}»", 25, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, Inches(1.24), Inches(4.78), Inches(3.1), Inches(0.34), "1 фраза вместо 5–10 экранов", 13, MIST, True, PP_ALIGN.CENTER)

    add_card(slide, Inches(5.22), Inches(2.58), Inches(3.25), Inches(3.02), CARD)
    add_text(slide, Inches(5.54), Inches(2.96), Inches(2.5), Inches(0.32), "Что делает агент", 16, accent, True)
    add_bullets(slide, Inches(5.56), Inches(3.58), Inches(2.36), Inches(1.42), actions, 13.3, WHITE, 0.07)

    add_card(slide, Inches(8.86), Inches(2.58), Inches(3.45), Inches(3.02), rgb("#102E33") if accent == MINT else CARD_2)
    add_text(slide, Inches(9.18), Inches(2.96), Inches(2.85), Inches(0.32), "Почему это продаёт", 16, accent, True)
    add_text(slide, Inches(9.22), Inches(3.66), Inches(2.55), Inches(0.84), result, 20, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, Inches(9.26), Inches(4.86), Inches(2.5), Inches(0.34), "меньше внимания к телефону — больше результата", 11.5, MIST, False, PP_ALIGN.CENTER)
    add_footer(slide, idx)


def make_deck():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # 1. Cover
    slide = prs.slides.add_slide(blank)
    add_bg(slide, "mint")
    add_pill(slide, Inches(0.76), Inches(0.58), "Premium v2", MINT, 1.55)
    add_text(slide, Inches(0.76), Inches(1.22), Inches(7.2), Inches(0.62), "AgentPhone", 50, WHITE, True)
    add_text(slide, Inches(0.8), Inches(2.08), Inches(7.0), Inches(0.94), "Личный AI-шник\nв кармане", 35, MINT, True)
    add_text(
        slide,
        Inches(0.83),
        Inches(3.34),
        Inches(6.4),
        Inches(0.88),
        "Не требует много внимания, но всегда под рукой: понимает цель, сам ведёт рутину и просит подтверждение только в важных местах.",
        17,
        MIST,
    )
    add_metric(slide, Inches(0.84), Inches(5.08), "1", "вход: цель")
    add_metric(slide, Inches(3.48), Inches(5.08), "5", "pocket-cases", BLUE)
    add_metric(slide, Inches(6.12), Inches(5.08), "L0/L1/L2", "контроль риска", GOLD)
    add_phone(slide, Inches(9.06), Inches(0.72), Inches(2.95), Inches(5.78), MINT)
    add_footer(slide, 1)

    # 2. Insight
    slide = prs.slides.add_slide(blank)
    add_bg(slide, "blue")
    add_title(slide, "Insight", "Смартфон стал работой", "Каждый день пользователь обслуживает интерфейсы: открывает, ищет, переключается, помнит, проверяет.", BLUE)
    words = [
        ("найти иконку", Inches(1.0), Inches(2.75), 21, MIST),
        ("открыть чат", Inches(3.65), Inches(3.25), 26, WHITE),
        ("вспомнить контекст", Inches(6.7), Inches(2.7), 21, MIST),
        ("заполнить форму", Inches(8.5), Inches(3.75), 24, WHITE),
        ("проверить оплату", Inches(2.25), Inches(4.62), 22, MIST),
        ("не забыть follow-up", Inches(5.35), Inches(5.1), 27, BLUE),
    ]
    for text, x, y, size, color in words:
        add_text(slide, x, y, Inches(3.2), Inches(0.4), text, size, color, True, PP_ALIGN.CENTER)
    add_card(slide, Inches(1.55), Inches(6.12), Inches(10.15), Inches(0.64), rgb("#111F35"))
    add_text(slide, Inches(1.9), Inches(6.29), Inches(9.45), Inches(0.22), "Проблема не в том, что ассистенты плохо отвечают. Проблема в том, что телефон всё ещё требует оператора.", 15, WHITE, True, PP_ALIGN.CENTER)
    add_footer(slide, 2)

    # 3. Big idea
    slide = prs.slides.add_slide(blank)
    add_bg(slide, "violet")
    add_title(slide, "Big idea", "Не ассистент поверх телефона. Телефон под управлением агента.", "Pixel становится руками и глазами AI, а приложения — временными инструментами для выполнения целей.", VIOLET)
    steps = [
        ("Цель", "«Сделай» вместо «Открой приложение»"),
        ("План", "агент сам выбирает шаги и инструменты"),
        ("Действие", "Accessibility, уведомления, сервисы, APK"),
        ("Память", "Personal KB хранит людей, места, привычки"),
    ]
    for i, (head, body) in enumerate(steps):
        x = Inches(0.96 + i * 3.05)
        add_card(slide, x, Inches(3.05), Inches(2.52), Inches(1.7), CARD_2)
        add_text(slide, x + Inches(0.22), Inches(3.32), Inches(2.0), Inches(0.32), head, 18, VIOLET if i == 0 else MINT, True, PP_ALIGN.CENTER)
        add_text(slide, x + Inches(0.32), Inches(3.98), Inches(1.78), Inches(0.46), body, 11.7, MIST, False, PP_ALIGN.CENTER)
    add_text(slide, Inches(1.15), Inches(5.64), Inches(11.0), Inches(0.46), "Ощущение для пользователя: рядом человек, который не мешает, но всегда может взять рутину на себя.", 21, MINT, True, PP_ALIGN.CENTER)
    add_footer(slide, 3)

    # 4. Five cases overview
    slide = prs.slides.add_slide(blank)
    add_bg(slide, "mint")
    add_title(slide, "5 pocket-cases", "Где AgentPhone продаёт себя сам", "Каждый кейс показывает не функцию, а снятую с пользователя микрорутину.", MINT)
    cases = [
        ("01", "Переписка", "согласует, отвечает, не забывает", MINT),
        ("02", "День", "следит за временем, дорогой и push", BLUE),
        ("03", "Сделки", "бронь, покупки, оплаты с L2", GOLD),
        ("04", "Телефон", "чистит, обновляет, настраивает", VIOLET),
        ("05", "Память", "follow-up, люди, обещания", ROSE),
    ]
    for i, (num, title, body, color) in enumerate(cases):
        add_case_card(slide, Inches(0.72 + i * 2.52), Inches(3.05), num, title, body, color)
    add_text(slide, Inches(1.18), Inches(5.78), Inches(10.9), Inches(0.38), "Формула: агент не требует внимания постоянно, но берёт управление в момент, когда человеку это нужно.", 18, WHITE, True, PP_ALIGN.CENTER)
    add_footer(slide, 4)

    add_case_slide(
        prs,
        5,
        "mint",
        MINT,
        "01",
        "Переписка без микроменеджмента",
        "Агент закрывает короткие коммуникации, где важны контекст и тон.",
        "Ответь маме, что буду через час",
        ["резолвит контакт", "выбирает канал", "пишет в привычном стиле", "просит approve перед отправкой"],
        "меньше забытых сообщений",
    )
    add_case_slide(
        prs,
        6,
        "blue",
        BLUE,
        "02",
        "День под контролем без лишних уведомлений",
        "Телефон сам следит за важным и не шумит, когда всё нормально.",
        "Скажи, когда подъедет такси",
        ["слушает push", "сверяет время и гео", "говорит вслух при триггере", "перепланирует при задержке"],
        "не нужно мониторить экран",
    )
    add_case_slide(
        prs,
        7,
        "gold",
        GOLD,
        "03",
        "Бронирования, покупки и оплаты",
        "Автопилот для рутины с безопасной остановкой на финальном шаге.",
        "Забронируй стол на завтра",
        ["ищет варианты", "заполняет формы", "ставит нужный APK", "деньги — только через L2"],
        "скорость без потери контроля",
    )
    add_case_slide(
        prs,
        8,
        "violet",
        VIOLET,
        "04",
        "Техподдержка телефона в кармане",
        "Личный администратор устройства, который держит телефон в порядке.",
        "Разберись, почему телефон тормозит",
        ["находит тяжёлые приложения", "чистит дубликаты", "обновляет нужное", "объясняет изменения"],
        "телефон не превращается в свалку",
    )
    add_case_slide(
        prs,
        9,
        "rose",
        ROSE,
        "05",
        "Память, follow-up и «не забыть»",
        "Агент держит договорённости и подсказывает следующий шаг без лишнего шума.",
        "Напомни, если я не отвечу инвестору",
        ["помнит контекст", "связывает людей и события", "создаёт follow-up", "предлагает следующий шаг"],
        "внешняя память без ухода",
    )

    # 10. How it works
    slide = prs.slides.add_slide(blank)
    add_bg(slide, "blue")
    add_title(slide, "How it works", "Observe → Plan → Act → Confirm → Remember", "Архитектура объясняется одним циклом, понятным покупателю и инвестору.", BLUE)
    cycle = [
        ("Observe", "голос, текст, push, календарь, гео"),
        ("Plan", "цель → шаги → инструменты"),
        ("Act", "клики, текст, APK, сервисы"),
        ("Confirm", "L0/L1/L2 по уровню риска"),
        ("Remember", "Personal KB и audit trail"),
    ]
    for i, (head, body) in enumerate(cycle):
        x = Inches(0.88 + i * 2.5)
        y = Inches(3.02)
        node = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x, y, Inches(1.0), Inches(1.0))
        node.fill.solid()
        node.fill.fore_color.rgb = BLUE if i % 2 else MINT
        node.line.fill.background()
        add_text(slide, x + Inches(0.06), y + Inches(0.34), Inches(0.88), Inches(0.2), str(i + 1), 20, INK, True, PP_ALIGN.CENTER)
        add_text(slide, x - Inches(0.28), Inches(4.34), Inches(1.6), Inches(0.28), head, 15, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, x - Inches(0.58), Inches(4.78), Inches(2.2), Inches(0.42), body, 10.5, MIST, False, PP_ALIGN.CENTER)
    add_text(slide, Inches(1.2), Inches(6.05), Inches(10.8), Inches(0.32), "Локально — приватное и простое. В облако — только сложное планирование обезличенным slice.", 15, BLUE, True, PP_ALIGN.CENTER)
    add_footer(slide, 10)

    # 11. Trust
    slide = prs.slides.add_slide(blank)
    add_bg(slide, "rose")
    add_title(slide, "Trust model", "Автономность без потери контроля", "Пользователь не должен доверять агенту вслепую: рискованные действия имеют понятный тормоз.", ROSE)
    rows = [
        ("L0", "Read-only", "агент действует молча", MINT),
        ("L1", "State change", "подтверждение в чате", GOLD),
        ("L2", "Деньги / третьи лица", "голосовое подтверждение", ROSE),
    ]
    for i, (lvl, head, body, color) in enumerate(rows):
        add_card(slide, Inches(1.0), Inches(2.72 + i * 1.02), Inches(5.35), Inches(0.72), CARD_2)
        add_text(slide, Inches(1.28), Inches(2.89 + i * 1.02), Inches(0.62), Inches(0.22), lvl, 17, color, True)
        add_text(slide, Inches(2.08), Inches(2.89 + i * 1.02), Inches(1.85), Inches(0.22), head, 14.5, WHITE, True)
        add_text(slide, Inches(4.02), Inches(2.91 + i * 1.02), Inches(1.8), Inches(0.2), body, 11.5, MIST)
    add_card(slide, Inches(7.1), Inches(2.72), Inches(4.95), Inches(2.78), rgb("#2E1830"))
    add_text(slide, Inches(7.48), Inches(3.08), Inches(4.2), Inches(0.28), "Sandbox для денег и памяти", 18, ROSE, True)
    add_bullets(slide, Inches(7.52), Inches(3.72), Inches(3.78), Inches(1.1), ["отдельная агентская карта", "лимиты и журнал операций", "E2EE backup, ключ у пользователя"], 14, WHITE, 0.08)
    add_footer(slide, 11)

    # 12. Differentiation
    slide = prs.slides.add_slide(blank)
    add_bg(slide, "violet")
    add_title(slide, "Differentiation", "Провалился не AI. Провалился лишний гаджет.", "Rabbit R1 и Humane AI Pin требовали новый предмет и новое поведение. AgentPhone живёт в привычном смартфоне.", VIOLET)
    add_card(slide, Inches(0.94), Inches(2.82), Inches(5.28), Inches(2.64), rgb("#381D33"))
    add_text(slide, Inches(1.34), Inches(3.18), Inches(4.45), Inches(0.34), "AI-gadgets", 22, ROSE, True, PP_ALIGN.CENTER)
    add_bullets(slide, Inches(1.55), Inches(3.92), Inches(3.8), Inches(0.9), ["ещё одна вещь", "ещё один интерфейс", "ограниченный доступ к жизни"], 17, WHITE, 0.08)
    add_card(slide, Inches(7.1), Inches(2.82), Inches(5.28), Inches(2.64), rgb("#122F36"))
    add_text(slide, Inches(7.5), Inches(3.18), Inches(4.45), Inches(0.34), "AgentPhone", 22, MINT, True, PP_ALIGN.CENTER)
    add_bullets(slide, Inches(7.72), Inches(3.92), Inches(3.8), Inches(0.9), ["тот же Pixel", "привычные приложения", "новый intelligence-layer"], 17, WHITE, 0.08)
    add_footer(slide, 12)

    # 13. Business
    slide = prs.slides.add_slide(blank)
    add_bg(slide, "gold")
    add_title(slide, "Business", "Продаём готовый AI-телефон, а не приложение", "Pixel + кастомный образ + agent-runtime + управляемая облачная инфраструктура.", GOLD)
    add_metric(slide, Inches(0.92), Inches(2.62), "Pixel", "прошитый под ключ", MINT)
    add_metric(slide, Inches(3.72), Inches(2.62), "BYOK", "контроль LLM-cost", BLUE)
    add_metric(slide, Inches(6.52), Inches(2.62), "$5–10", "лимит по умолчанию", GOLD)
    add_metric(slide, Inches(9.32), Inches(2.62), "100/мес", "PMF-метрика", ROSE)
    add_card(slide, Inches(1.35), Inches(4.55), Inches(10.6), Inches(0.88), rgb("#2D2A1A"))
    add_text(slide, Inches(1.72), Inches(4.8), Inches(9.9), Inches(0.26), "Open-source — канал доверия и distribution. Доход — устройство и agent-инфраструктура.", 16, WHITE, True, PP_ALIGN.CENTER)
    add_footer(slide, 13)

    # 14. Ask
    slide = prs.slides.add_slide(blank)
    add_bg(slide, "mint")
    add_title(slide, "Ask", "Первые 100 устройств должны доказать новую привычку", "Покупают не AI-фичу. Покупают ощущение, что телефон наконец-то работает на человека.", MINT)
    columns = [
        ("Доказать", ["100 устройств/мес", "5 killer-cases", "удержание через Personal KB"], MINT),
        ("Нужно", ["ранние покупатели", "Pixel-поставка", "ресурсы на runtime/OTA"], GOLD),
        ("Верить", ["привычный форм-фактор", "local-first privacy", "контроль риска L0/L1/L2"], BLUE),
    ]
    for i, (head, items, color) in enumerate(columns):
        x = Inches(0.95 + i * 4.08)
        add_card(slide, x, Inches(2.72), Inches(3.5), Inches(2.45), CARD_2)
        add_text(slide, x + Inches(0.32), Inches(3.08), Inches(2.85), Inches(0.32), head, 19, color, True, PP_ALIGN.CENTER)
        add_bullets(slide, x + Inches(0.62), Inches(3.78), Inches(2.25), Inches(0.88), items, 14, WHITE, 0.08)
    add_text(slide, Inches(1.0), Inches(6.15), Inches(11.3), Inches(0.5), "AgentPhone: smartphone becomes your pocket agent.", 27, MINT, True, PP_ALIGN.CENTER)
    add_footer(slide, 14)

    prs.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    print(make_deck())
