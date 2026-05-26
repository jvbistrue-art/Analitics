from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT_DIR = Path(__file__).resolve().parent
OUT_FILE = OUT_DIR / "AgentPhone_Sales_Deck.pptx"

WIDE_W = Inches(13.333)
WIDE_H = Inches(7.5)

BG = RGBColor(10, 14, 24)
PANEL = RGBColor(18, 25, 42)
PANEL_2 = RGBColor(23, 35, 57)
TEXT = RGBColor(244, 247, 251)
MUTED = RGBColor(162, 174, 195)
ACCENT = RGBColor(82, 236, 171)
ACCENT_2 = RGBColor(94, 156, 255)
WARNING = RGBColor(255, 198, 93)
PINK = RGBColor(255, 114, 147)

FONT = "Aptos"


def add_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG

    # Ambient glow blocks: lightweight visual identity without external assets.
    glow_1 = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL, Inches(9.8), Inches(-0.7), Inches(4.4), Inches(4.4)
    )
    glow_1.fill.solid()
    glow_1.fill.fore_color.rgb = RGBColor(20, 77, 91)
    glow_1.fill.transparency = 35
    glow_1.line.fill.background()

    glow_2 = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL, Inches(-1.5), Inches(5.8), Inches(4.2), Inches(4.2)
    )
    glow_2.fill.solid()
    glow_2.fill.fore_color.rgb = RGBColor(53, 39, 92)
    glow_2.fill.transparency = 42
    glow_2.line.fill.background()


def add_textbox(slide, x, y, w, h, text="", font_size=24, color=TEXT, bold=False, align=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=None, kicker=None):
    if kicker:
        add_textbox(slide, Inches(0.75), Inches(0.43), Inches(11.8), Inches(0.28), kicker.upper(), 10, ACCENT, True)
    add_textbox(slide, Inches(0.75), Inches(0.86), Inches(11.8), Inches(0.9), title, 34, TEXT, True)
    if subtitle:
        add_textbox(slide, Inches(0.78), Inches(1.72), Inches(10.9), Inches(0.58), subtitle, 17, MUTED)


def add_footer(slide, idx):
    add_textbox(slide, Inches(0.75), Inches(7.02), Inches(2.4), Inches(0.25), "AgentPhone", 9, MUTED, True)
    add_textbox(slide, Inches(12.08), Inches(7.02), Inches(0.5), Inches(0.25), f"{idx:02}", 9, MUTED, True, PP_ALIGN.RIGHT)


def add_panel(slide, x, y, w, h, color=PANEL, transparency=0, radius=True):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    panel = slide.shapes.add_shape(shape_type, x, y, w, h)
    panel.fill.solid()
    panel.fill.fore_color.rgb = color
    panel.fill.transparency = transparency
    panel.line.color.rgb = RGBColor(44, 57, 82)
    panel.line.transparency = 45
    return panel


def add_chip(slide, x, y, label, color=ACCENT, w=1.8):
    chip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, Inches(w), Inches(0.36))
    chip.fill.solid()
    chip.fill.fore_color.rgb = color
    chip.fill.transparency = 12
    chip.line.fill.background()
    tf = chip.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = FONT
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = BG
    return chip


def add_bullets(slide, x, y, w, h, bullets, font_size=18, color=TEXT, gap=0.16):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.space_after = Pt(gap * 72)
        p.font.name = FONT
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p._p.get_or_add_pPr().set("marL", "0")
    return box


def add_metric(slide, x, y, w, h, number, label, accent=ACCENT):
    add_panel(slide, x, y, w, h, PANEL_2)
    add_textbox(slide, x + Inches(0.22), y + Inches(0.2), w - Inches(0.44), Inches(0.48), number, 28, accent, True)
    add_textbox(slide, x + Inches(0.24), y + Inches(0.82), w - Inches(0.48), h - Inches(0.92), label, 12, MUTED)


def add_phone_mock(slide, x, y, w, h, title="AgentPhone"):
    phone = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    phone.fill.solid()
    phone.fill.fore_color.rgb = RGBColor(8, 11, 18)
    phone.line.color.rgb = RGBColor(71, 91, 126)
    phone.line.width = Pt(1.4)

    screen = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        x + Inches(0.15),
        y + Inches(0.22),
        w - Inches(0.30),
        h - Inches(0.44),
    )
    screen.fill.solid()
    screen.fill.fore_color.rgb = RGBColor(13, 20, 35)
    screen.line.fill.background()

    add_textbox(slide, x + Inches(0.35), y + Inches(0.46), w - Inches(0.7), Inches(0.28), title, 11, MUTED, True, PP_ALIGN.CENTER)
    add_chat_bubble(slide, x + Inches(0.38), y + Inches(1.03), w - Inches(0.76), "Хочу попить пива завтра с Васей", False)
    add_chat_bubble(slide, x + Inches(0.38), y + Inches(1.78), w - Inches(0.76), "Нашёл 3 немецких бара. Написать Васе и забронировать?", True)
    add_chat_bubble(slide, x + Inches(0.38), y + Inches(2.78), w - Inches(0.76), "Да, на 19:30", False)
    add_chat_bubble(slide, x + Inches(0.38), y + Inches(3.48), w - Inches(0.76), "Готово: Вася подтвердил, стол забронирован, календарь обновлён.", True)

    mic = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x + w / 2 - Inches(0.32), y + h - Inches(0.82), Inches(0.64), Inches(0.64))
    mic.fill.solid()
    mic.fill.fore_color.rgb = ACCENT
    mic.line.fill.background()
    add_textbox(slide, x + w / 2 - Inches(0.08), y + h - Inches(0.68), Inches(0.16), Inches(0.18), "•", 21, BG, True, PP_ALIGN.CENTER)


def add_chat_bubble(slide, x, y, w, text, agent):
    bubble_w = w * 0.82
    bx = x if agent else x + w - bubble_w
    color = RGBColor(30, 45, 72) if agent else ACCENT
    txt_color = TEXT if agent else BG
    bubble = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, bx, y, bubble_w, Inches(0.58))
    bubble.fill.solid()
    bubble.fill.fore_color.rgb = color
    bubble.line.fill.background()
    tf = bubble.text_frame
    tf.clear()
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(9.5)
    r.font.color.rgb = txt_color


def add_arrow(slide, x1, y1, x2, y2, color=ACCENT):
    line = slide.shapes.add_connector(1, x1, y1, x2, y2)
    line.line.color.rgb = color
    line.line.width = Pt(1.6)
    line.line.transparency = 8
    return line


def make_deck():
    prs = Presentation()
    prs.slide_width = WIDE_W
    prs.slide_height = WIDE_H
    blank = prs.slide_layouts[6]

    # 1. Cover
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_chip(slide, Inches(0.75), Inches(0.55), "AI-native smartphone", ACCENT, 2.2)
    add_textbox(slide, Inches(0.75), Inches(1.34), Inches(7.0), Inches(1.1), "AgentPhone", 48, TEXT, True)
    add_textbox(slide, Inches(0.78), Inches(2.34), Inches(7.0), Inches(0.9), "Не «AI в телефоне».\nТелефон под управлением AI.", 26, ACCENT, True)
    add_textbox(
        slide,
        Inches(0.78),
        Inches(3.52),
        Inches(6.55),
        Inches(1.02),
        "Личный agent-помощник, у которого Pixel — это руки, глаза и безопасный доступ к цифровой жизни пользователя.",
        18,
        MUTED,
    )
    add_metric(slide, Inches(0.78), Inches(5.12), Inches(1.85), Inches(1.08), "30+", "приложений сегодня управляются человеком")
    add_metric(slide, Inches(2.9), Inches(5.12), Inches(1.85), Inches(1.08), "1", "чат-интерфейс вместо иконок")
    add_metric(slide, Inches(5.02), Inches(5.12), Inches(1.85), Inches(1.08), "100", "устройств/мес — первая PMF-метрика")
    add_phone_mock(slide, Inches(8.62), Inches(0.7), Inches(3.1), Inches(5.95))
    add_footer(slide, 1)

    # 2. Problem
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Проблема: смартфон стал пультом от 30+ сервисов", "Пользователь не хочет управлять интерфейсами. Он хочет, чтобы задача была сделана.", "Problem")
    add_panel(slide, Inches(0.75), Inches(2.55), Inches(3.72), Inches(3.28))
    add_textbox(slide, Inches(1.05), Inches(2.88), Inches(3.05), Inches(0.4), "Сегодня", 18, WARNING, True)
    add_bullets(slide, Inches(1.05), Inches(3.46), Inches(3.05), Inches(1.9), ["помнить, где какая иконка", "кликать 5–10 экранов", "держать контекст в голове"], 18)
    add_panel(slide, Inches(4.82), Inches(2.55), Inches(3.72), Inches(3.28))
    add_textbox(slide, Inches(5.12), Inches(2.88), Inches(3.05), Inches(0.4), "Ассистенты", 18, PINK, True)
    add_bullets(slide, Inches(5.12), Inches(3.46), Inches(3.05), Inches(1.9), ["заменяют клики голосом", "не планируют end-to-end", "не владеют состоянием задачи"], 18)
    add_panel(slide, Inches(8.89), Inches(2.55), Inches(3.72), Inches(3.28))
    add_textbox(slide, Inches(9.19), Inches(2.88), Inches(3.05), Inches(0.4), "Рынок ждёт", 18, ACCENT, True)
    add_bullets(slide, Inches(9.19), Inches(3.46), Inches(3.05), Inches(1.9), ["один вход: цель", "память о человеке", "действия в приложениях"], 18)
    add_footer(slide, 2)

    # 3. Solution
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Решение: launcher превращается в agent-runtime", "Пользователь формулирует цель. AgentPhone планирует, действует, проверяет результат и просит подтверждение только там, где есть риск.", "Solution")
    add_phone_mock(slide, Inches(0.86), Inches(2.3), Inches(2.65), Inches(4.55), "Chat Launcher")
    steps = [
        ("1", "Понять цель", "распознать intent и недостающие параметры"),
        ("2", "Разложить на план", "выбрать приложение, системный сервис или cloud-agent"),
        ("3", "Сделать руками Pixel", "Accessibility: пишет, кликает, навигирует"),
        ("4", "Проверить и запомнить", "мониторинг, replanning, Personal KB"),
    ]
    for i, (num, head, body) in enumerate(steps):
        x = Inches(4.12 + (i % 2) * 4.12)
        y = Inches(2.38 + (i // 2) * 1.76)
        add_panel(slide, x, y, Inches(3.58), Inches(1.24), PANEL_2)
        add_textbox(slide, x + Inches(0.24), y + Inches(0.22), Inches(0.42), Inches(0.4), num, 22, ACCENT, True)
        add_textbox(slide, x + Inches(0.78), y + Inches(0.19), Inches(2.45), Inches(0.32), head, 16, TEXT, True)
        add_textbox(slide, x + Inches(0.78), y + Inches(0.58), Inches(2.4), Inches(0.4), body, 11.5, MUTED)
    add_textbox(slide, Inches(4.15), Inches(6.0), Inches(7.8), Inches(0.56), "Метафора: Jarvis из Iron Man, не Google Assistant.", 22, ACCENT, True, PP_ALIGN.CENTER)
    add_footer(slide, 3)

    # 4. Hero use case
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Флагманский сценарий показывает магию продукта", "«Хочу попить пива с Васей в немецком баре завтра» — один запрос вместо цепочки из мессенджера, карт, браузера, бронирования и календаря.", "Hero demo")
    labels = [
        ("Понял", "Вася = контакт, завтра = дата, немецкий бар = предпочтение"),
        ("Нашёл", "3 бара рядом, проверил рейтинг и доступность"),
        ("Согласовал", "написал Васе в предпочитаемом канале"),
        ("Забронировал", "поставил APK бара при необходимости"),
        ("Зафиксировал", "календарь, напоминание, replan при изменениях"),
    ]
    x0 = Inches(0.92)
    for i, (head, body) in enumerate(labels):
        x = x0 + Inches(i * 2.42)
        y = Inches(3.08)
        node = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x, y, Inches(0.62), Inches(0.62))
        node.fill.solid()
        node.fill.fore_color.rgb = ACCENT if i == 0 else ACCENT_2
        node.line.fill.background()
        add_textbox(slide, x + Inches(0.18), y + Inches(0.12), Inches(0.25), Inches(0.22), str(i + 1), 14, BG, True, PP_ALIGN.CENTER)
        if i < len(labels) - 1:
            add_arrow(slide, x + Inches(0.72), y + Inches(0.31), x + Inches(2.08), y + Inches(0.31), ACCENT_2)
        add_textbox(slide, x - Inches(0.12), y + Inches(0.86), Inches(1.42), Inches(0.34), head, 15, TEXT, True, PP_ALIGN.CENTER)
        add_textbox(slide, x - Inches(0.42), y + Inches(1.3), Inches(2.05), Inches(0.72), body, 10.5, MUTED, False, PP_ALIGN.CENTER)
    add_panel(slide, Inches(1.1), Inches(5.86), Inches(11.1), Inches(0.74), RGBColor(17, 43, 44))
    add_textbox(slide, Inches(1.46), Inches(6.05), Inches(10.4), Inches(0.3), "Продажа не в «умном чате», а в выполненной задаче: пользователь получает результат без ручного управления приложениями.", 16, ACCENT, True, PP_ALIGN.CENTER)
    add_footer(slide, 4)

    # 5. Product pillars
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Четыре обещания продукта", "Они переводят техническую архитектуру в понятную покупателю ценность.", "Value proposition")
    pillars = [
        ("Думает", "декомпозирует цель, планирует шаги, перепланирует при изменениях", ACCENT),
        ("Действует", "использует приложения как инструменты: пишет, кликает, ставит APK", ACCENT_2),
        ("Помнит", "Personal KB: контакты, события, привычки — не в приложениях, а у пользователя", WARNING),
        ("Защищает", "L0/L1/L2 подтверждения, sandbox-карта, E2EE backup, audit trail", PINK),
    ]
    for i, (head, body, color) in enumerate(pillars):
        x = Inches(0.82 + (i % 2) * 6.1)
        y = Inches(2.46 + (i // 2) * 1.76)
        add_panel(slide, x, y, Inches(5.55), Inches(1.2), PANEL_2)
        add_chip(slide, x + Inches(0.26), y + Inches(0.2), head, color, 1.25)
        add_textbox(slide, x + Inches(1.75), y + Inches(0.25), Inches(3.45), Inches(0.58), body, 15, TEXT, True)
    add_footer(slide, 5)

    # 6. Trust and safety
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Доверие встроено в UX, а не добавлено после", "AgentPhone просит подтверждение по уровню риска и оставляет проверяемый журнал действий.", "Safety")
    matrix = [
        ("L0", "Read-only", "агент действует молча", ACCENT),
        ("L1", "State change", "подтверждение в чате", WARNING),
        ("L2", "Деньги / третьи лица", "голосовое подтверждение", PINK),
    ]
    for i, (lvl, name, body, color) in enumerate(matrix):
        y = Inches(2.45 + i * 1.16)
        add_panel(slide, Inches(0.9), y, Inches(5.5), Inches(0.86), PANEL_2)
        add_textbox(slide, Inches(1.2), y + Inches(0.16), Inches(0.62), Inches(0.3), lvl, 19, color, True)
        add_textbox(slide, Inches(2.0), y + Inches(0.16), Inches(1.75), Inches(0.3), name, 16, TEXT, True)
        add_textbox(slide, Inches(3.86), y + Inches(0.19), Inches(2.05), Inches(0.3), body, 12, MUTED)
    add_panel(slide, Inches(7.0), Inches(2.38), Inches(5.42), Inches(3.56), RGBColor(28, 28, 49))
    add_textbox(slide, Inches(7.38), Inches(2.75), Inches(4.8), Inches(0.36), "Финансовый sandbox", 19, ACCENT, True)
    add_bullets(
        slide,
        Inches(7.38),
        Inches(3.36),
        Inches(4.35),
        Inches(1.82),
        ["отдельная агентская карта", "лимиты и журнал операций", "компрометация агента ≠ компрометация всех средств"],
        16,
    )
    add_textbox(slide, Inches(7.38), Inches(5.35), Inches(4.55), Inches(0.32), "Память: локально + E2EE backup, ключ у пользователя.", 13, MUTED, True)
    add_footer(slide, 6)

    # 7. Architecture
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Архитектура: local-first там, где важна приватность", "Сложное планирование уходит в облако только обезличенным slice. Простые и чувствительные задачи остаются на устройстве.", "Architecture")
    columns = [
        ("Pixel + LineageOS", ["чат-launcher", "Agent Runtime", "AccessibilityService", "NotificationListener", "PackageInstaller"], ACCENT),
        ("Local AI", ["LLM Tiny 1B router", "LLM Small 3–4B", "Whisper-small ASR", "Embeddings + Vector DB", "Silero TTS optional"], ACCENT_2),
        ("Cloud", ["Planner / Researcher", "LLM API via BYOK", "OTA origin + CDN", "E2EE backup storage", "monitoring agents"], WARNING),
    ]
    for i, (head, items, color) in enumerate(columns):
        x = Inches(0.86 + i * 4.15)
        add_panel(slide, x, Inches(2.42), Inches(3.55), Inches(3.68), PANEL_2)
        add_textbox(slide, x + Inches(0.24), Inches(2.78), Inches(3.0), Inches(0.38), head, 17, color, True)
        add_bullets(slide, x + Inches(0.34), Inches(3.42), Inches(2.85), Inches(2.0), items, 13.5, TEXT, 0.08)
    add_textbox(slide, Inches(0.88), Inches(6.42), Inches(11.7), Inches(0.32), "Size budget моделей: ~5 GB. RAM peak: ~3–4 GB. OTA через стандартный Android update_engine с A/B rollback.", 13.5, MUTED, True, PP_ALIGN.CENTER)
    add_footer(slide, 7)

    # 8. Market timing
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Почему сейчас", "Несколько технологий впервые стали достаточно зрелыми, чтобы AI мог управлять телефоном, а не просто отвечать в чате.", "Timing")
    reasons = [
        ("On-device LLM", "1–4B модели уже помещаются в мобильный budget для роутинга и приватных задач"),
        ("Android APIs", "Accessibility, уведомления, календарь, контакты и A/B OTA дают агенту реальные руки"),
        ("Cloud agents", "сложные задачи можно выносить в облако без раскрытия всей персональной памяти"),
        ("Open-source trust", "LineageOS + Apache 2.0 снижают барьер доверия для ранних пользователей"),
    ]
    for i, (head, body) in enumerate(reasons):
        x = Inches(0.94 + (i % 2) * 6.03)
        y = Inches(2.38 + (i // 2) * 1.72)
        add_panel(slide, x, y, Inches(5.35), Inches(1.14), PANEL_2)
        add_textbox(slide, x + Inches(0.3), y + Inches(0.23), Inches(2.0), Inches(0.3), head, 16, ACCENT, True)
        add_textbox(slide, x + Inches(2.34), y + Inches(0.21), Inches(2.65), Inches(0.48), body, 11.5, TEXT)
    add_footer(slide, 8)

    # 9. Business model
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Бизнес-модель: продаём готовый AI-телефон", "Open-source — канал доверия и distribution. Деньги — в устройстве под ключ и управляемой agent-инфраструктуре.", "Business")
    add_metric(slide, Inches(0.92), Inches(2.5), Inches(2.55), Inches(1.25), "Pixel + OS", "прошитые телефоны под ключ")
    add_metric(slide, Inches(3.78), Inches(2.5), Inches(2.55), Inches(1.25), "$5–10", "default cloud budget на пользователя в месяц", ACCENT_2)
    add_metric(slide, Inches(6.64), Inches(2.5), Inches(2.55), Inches(1.25), "BYOK", "контроль cost и выбор LLM-провайдера", WARNING)
    add_metric(slide, Inches(9.5), Inches(2.5), Inches(2.55), Inches(1.25), "100/мес", "первая PMF-метрика: проданные устройства", PINK)
    add_panel(slide, Inches(1.2), Inches(4.55), Inches(10.9), Inches(1.28), RGBColor(20, 42, 54))
    add_textbox(slide, Inches(1.58), Inches(4.86), Inches(10.1), Inches(0.38), "Покупатель платит не за ещё один ассистент, а за смартфон, который перестаёт требовать оператора.", 20, ACCENT, True, PP_ALIGN.CENTER)
    add_footer(slide, 9)

    # 10. Competitive contrast
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Почему AgentPhone ≠ Rabbit R1 и Humane AI Pin", "Они пытались добавить новый гаджет. AgentPhone заменяет основной смартфон, сохраняя привычный форм-фактор и приложения.", "Competition")
    add_panel(slide, Inches(0.92), Inches(2.42), Inches(5.35), Inches(3.25), RGBColor(48, 24, 38))
    add_textbox(slide, Inches(1.3), Inches(2.8), Inches(4.5), Inches(0.4), "AI-gadgets", 22, PINK, True, PP_ALIGN.CENTER)
    add_bullets(slide, Inches(1.4), Inches(3.55), Inches(4.2), Inches(1.42), ["ещё одна вещь в кармане", "нужно менять поведение", "ограниченный доступ к приложениям"], 18)
    add_panel(slide, Inches(7.08), Inches(2.42), Inches(5.35), Inches(3.25), RGBColor(17, 45, 43))
    add_textbox(slide, Inches(7.46), Inches(2.8), Inches(4.5), Inches(0.4), "AgentPhone", 22, ACCENT, True, PP_ALIGN.CENTER)
    add_bullets(slide, Inches(7.56), Inches(3.55), Inches(4.2), Inches(1.42), ["тот же Pixel в кармане", "привычные жесты и приложения", "новый intelligence-layer внутри"], 18)
    add_footer(slide, 10)

    # 11. Go-to-market
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Первый рынок: люди, для которых время дороже настройки", "Продукт продаётся через сильные сценарии, а не через список фич.", "Go-to-market")
    segments = [
        ("Founders / operators", "много переписок, встреч, оплат, быстрых решений"),
        ("Tech early adopters", "готовы платить за новый интерфейс к жизни"),
        ("Power users Android", "ценят кастомизацию, privacy и open-source"),
    ]
    for i, (head, body) in enumerate(segments):
        x = Inches(0.92 + i * 4.08)
        add_panel(slide, x, Inches(2.65), Inches(3.5), Inches(2.08), PANEL_2)
        add_textbox(slide, x + Inches(0.26), Inches(3.0), Inches(3.0), Inches(0.46), head, 17, ACCENT, True, PP_ALIGN.CENTER)
        add_textbox(slide, x + Inches(0.42), Inches(3.72), Inches(2.65), Inches(0.7), body, 13, TEXT, False, PP_ALIGN.CENTER)
    add_panel(slide, Inches(1.55), Inches(5.42), Inches(10.25), Inches(0.82), RGBColor(24, 34, 58))
    add_textbox(slide, Inches(1.94), Inches(5.65), Inches(9.5), Inches(0.26), "Стартовая витрина: 5 эталонных user stories, где агент делает задачу end-to-end.", 16, MUTED, True, PP_ALIGN.CENTER)
    add_footer(slide, 11)

    # 12. Ask
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_title(slide, "Следующий шаг: превратить прототип в первые 100 устройств", "Фокус — доказать, что люди покупают не «AI-фичу», а новый способ пользоваться смартфоном.", "Ask")
    add_panel(slide, Inches(0.95), Inches(2.35), Inches(3.7), Inches(3.2), PANEL_2)
    add_textbox(slide, Inches(1.28), Inches(2.75), Inches(3.0), Inches(0.36), "Что доказываем", 19, ACCENT, True)
    add_bullets(slide, Inches(1.28), Inches(3.38), Inches(2.9), Inches(1.5), ["100 проданных устройств/мес", "5 killer-сценариев работают", "удержание через Personal KB"], 16)
    add_panel(slide, Inches(4.88), Inches(2.35), Inches(3.7), Inches(3.2), PANEL_2)
    add_textbox(slide, Inches(5.21), Inches(2.75), Inches(3.0), Inches(0.36), "Что нужно", 19, WARNING, True)
    add_bullets(slide, Inches(5.21), Inches(3.38), Inches(2.9), Inches(1.5), ["ранние покупатели", "партнёры по Pixel-поставке", "ресурсы на agent-runtime и OTA"], 16)
    add_panel(slide, Inches(8.81), Inches(2.35), Inches(3.7), Inches(3.2), PANEL_2)
    add_textbox(slide, Inches(9.14), Inches(2.75), Inches(3.0), Inches(0.36), "Почему верить", 19, ACCENT_2, True)
    add_bullets(slide, Inches(9.14), Inches(3.38), Inches(2.9), Inches(1.5), ["существующий Android-форм-фактор", "local-first безопасность", "open-source доверие"], 16)
    add_textbox(slide, Inches(1.1), Inches(6.25), Inches(11.1), Inches(0.42), "AgentPhone: smartphone becomes an agent.", 26, ACCENT, True, PP_ALIGN.CENTER)
    add_footer(slide, 12)

    prs.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    path = make_deck()
    print(path)
